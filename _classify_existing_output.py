"""对已抓取的结果表补跑「标签 + 情感分析」，不重新抓取。

用法:
  python _classify_existing_output.py <输入xlsx> <输出xlsx> [并发数]

复刻 main.py 的写入格式：
  - 「评论+标签+情感」：一条评论命中 N 个标签 -> 展开成 N 行，共享同一序号
  - 「正文内容+标签+情感」：正文命中 N 个标签 -> 展开成 N 行，共享同一序号
"""
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor

import openpyxl
from openpyxl.styles import Alignment

from classifier import classify_content_expanded

IN = sys.argv[1] if len(sys.argv) > 1 else "dist/output/scrape_result_20260911_124803.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "dist/output/scrape_result_20260911_124803_classified.xlsx"
WORKERS = int(sys.argv[3]) if len(sys.argv) > 3 else 8

_lock = threading.Lock()
_done = [0]


def classify_one(text: str, is_post: bool = False):
    try:
        return classify_content_expanded(text, is_post_content=is_post)
    except Exception as e:
        with _lock:
            print(f"  [warn] 分类失败: {type(e).__name__}: {e}", flush=True)
        return []


def main():
    t0 = time.time()
    wb = openpyxl.load_workbook(IN)

    # ── Sheet2: 评论 + 标签 + 情感 ─────────────────────────────
    ws = wb["评论+标签+情感"]
    old = list(ws.iter_rows(min_row=2, values_only=True))
    old = [r for r in old if r and r[4] is not None]
    print(f"[i] 待分类评论 {len(old)} 条，并发 {WORKERS}", flush=True)

    contents = [r[4] for r in old]
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        results = []
        futs = [ex.submit(classify_one, c) for c in contents]
        for i, fu in enumerate(futs, 1):
            results.append(fu.result())
            with _lock:
                if i % 50 == 0 or i == len(futs):
                    print(f"    评论分类 {i}/{len(futs)}  用时 {time.time()-t0:.0f}s", flush=True)

    # 重建 sheet
    del wb["评论+标签+情感"]
    ws2 = wb.create_sheet("评论+标签+情感", 1)
    ws2.append(["序号", "情感偏向", "标签", "是否与其他酒店比较", "内容", "平台"])
    total_rows = 0
    for (seq, _sent, _tag, _cmp, content, platform), expanded in zip(old, results):
        if expanded:
            for item in expanded:
                ws2.append([seq, item["sentiment"], item["tag"], item["has_comparison"], content, platform])
                total_rows += 1
        else:
            ws2.append([seq, "中性", "/", "否", content, platform])
            total_rows += 1
    print(f"[i] 评论表写入 {total_rows} 行（{len(old)} 条评论）", flush=True)

    # ── Sheet3: 正文 + 标签 + 情感 ─────────────────────────────
    if "正文内容+标签+情感" in wb.sheetnames:
        ws3old = wb["正文内容+标签+情感"]
        rows3 = [r for r in ws3old.iter_rows(min_row=2, values_only=True) if r and r[3]]
        del wb["正文内容+标签+情感"]
        ws3 = wb.create_sheet("正文内容+标签+情感")
        ws3.append(["序号", "标题", "平台", "内容摘要", "情感偏向", "标签", "是否与其他酒店比较", "原文/评论链接"])
        if rows3:
            texts = []
            for r in rows3:
                disp = str(r[3] or "")
                body = disp.split("\n")[0]
                if body.startswith("正文："):
                    body = body[len("正文："):]
                texts.append(body)
            with ThreadPoolExecutor(max_workers=min(WORKERS, len(texts))) as ex:
                res3 = list(ex.map(lambda t: classify_one(t, is_post=True), texts))
            for r, expanded in zip(rows3, res3):
                seq, title, platform, disp, _s, _t, _c, url = r
                if expanded:
                    for item in expanded:
                        ws3.append([seq, title, platform, disp, item["sentiment"],
                                    item["tag"], item["has_comparison"], url])
                else:
                    ws3.append([seq, title, platform, disp, "/", "/", "否", url])
            print(f"[i] 正文分类完成: {[e['tag']+'/'+e['sentiment'] for e in res3[0]] if res3 and res3[0] else '无标签'}",
                  flush=True)

    # 统一列宽/换行，跟 main.py 观感一致
    for name in ("评论+标签+情感", "正文内容+标签+情感"):
        if name in wb.sheetnames:
            s = wb[name]
            s.column_dimensions["E" if name.startswith("评论") else "D"].width = 80
            for row in s.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)
    print(f"\n>>> 已保存 {OUT}  总耗时 {time.time()-t0:.0f}s", flush=True)


if __name__ == "__main__":
    main()
