# -*- coding: utf-8 -*-
"""Re-classify comments & 正文 in scrape_result_20260807_103801.xlsx and write a new file.

落地格式：一标签一行（与 main.py 一致）。多标签 → 同一序号重复多行。
空结果保持 main.py 的占位：评论「中性//否」、正文「///否」。
"""
import logging
import os
import re
import sys
import time

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
logging.basicConfig(level=logging.ERROR)

from classifier import classify_content

SRC = r"C:\Users\chenwentao\Desktop\私人设计\scrape_result_20260807_103801.xlsx"
DST = r"C:\Users\chenwentao\Desktop\私人设计\scrape_result_20260807_103801_重新标注.xlsx"

# 口播垃圾标记：这类口播不是可用口播内容，分类时退回正文
_GARBAGE_MARKERS = [
    "没有旁白", "没有口播", "无旁白", "无口播", "歌曲演唱",
    "The user wants", "no actual speech", "no words are present",
    "视频呈现", "视频内容",
]
_NON_CJK = re.compile(r"[^一-鿿]")


def is_garbage_transcription(text: str) -> bool:
    t = text.strip()
    if not t:
        return True
    if any(m in t for m in _GARBAGE_MARKERS):
        return True
    # 非中文占比过高（如法文歌词、英文拒转泄漏）→ 视为垃圾口播
    if len(_NON_CJK.findall(t)) / len(t) > 0.4:
        return True
    return False


def pick_classify_text(content_display: str) -> str:
    """按 main.py 逻辑重建 classify_text：有可用口播用口播，否则用正文。"""
    post_content = content_display.strip()
    transcription = ""
    if "口播：" in content_display:
        head, _, tail = content_display.partition("口播：")
        post_content = head.strip()
        transcription = tail.strip()
    if post_content.startswith("正文："):
        post_content = post_content[len("正文："):].strip()
    if transcription and not is_garbage_transcription(transcription):
        return transcription
    return post_content


def classify_expanded(text: str, is_post: bool) -> list[dict]:
    raw = classify_content(text, retries=1, is_post_content=is_post)
    expanded = []
    for tag_str in raw["tags"]:
        parts = tag_str.rsplit("-", 1)
        if len(parts) == 2:
            expanded.append({
                "tag": parts[0],
                "sentiment": parts[1],
                "has_comparison": raw["has_comparison"],
            })
    return expanded


def main():
    print(f"读取: {SRC}", flush=True)
    wb = openpyxl.load_workbook(SRC)

    ws2 = wb["评论+标签+情感"]
    ws3 = wb["正文内容+标签+情感"]

    # ── 读取原始数据 ──
    comment_rows = [[c.value for c in row] for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row)]
    post_rows = [[c.value for c in row] for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row)]
    print(f"评论行: {len(comment_rows)}  正文行: {len(post_rows)}", flush=True)

    # 去重内容，避免重复 API 调用
    comment_cache: dict[str, list[dict]] = {}
    for r in comment_rows:
        c = str(r[4]).strip() if r[4] else ""
        if c and c not in comment_cache:
            comment_cache[c] = None
    post_cache: dict[str, list[dict]] = {}
    for r in post_rows:
        c = str(r[3]).strip() if r[3] else ""
        if c and c not in post_cache:
            post_cache[c] = None
    print(f"去重后评论: {len(comment_cache)}  正文: {len(post_cache)}", flush=True)

    # ── 分类评论 ──
    print(f"\n{'='*60}\n评论分类 ({len(comment_cache)} 条)\n{'='*60}", flush=True)
    t0 = time.time()
    done = 0
    for content in comment_cache:
        expanded = classify_expanded(content, is_post=False)
        comment_cache[content] = expanded
        done += 1
        if done % 50 == 0:
            el = time.time() - t0
            print(f"  评论 {done}/{len(comment_cache)}  {el:.0f}s", flush=True)
    print(f"  评论完成，耗时 {time.time()-t0:.0f}s", flush=True)

    # ── 分类正文 ──
    print(f"\n{'='*60}\n正文分类 ({len(post_cache)} 条)\n{'='*60}", flush=True)
    t0 = time.time()
    done = 0
    for content in post_cache:
        classify_text = pick_classify_text(content)
        expanded = classify_expanded(classify_text, is_post=True) if classify_text.strip() else []
        post_cache[content] = expanded
        done += 1
        if done % 20 == 0:
            el = time.time() - t0
            print(f"  正文 {done}/{len(post_cache)}  {el:.0f}s", flush=True)
    print(f"  正文完成，耗时 {time.time()-t0:.0f}s", flush=True)

    # ── 重建 sheet2：一标签一行 ──
    ws2.delete_rows(2, ws2.max_row - 1)
    c_lab = c_empty = 0
    for r in comment_rows:
        seq, content = r[0], str(r[4]).strip() if r[4] else ""
        platform = r[5]
        expanded = comment_cache.get(content, [])
        if expanded:
            for item in expanded:
                ws2.append([seq, item["sentiment"], item["tag"], item["has_comparison"], content, platform])
            c_lab += 1
        else:
            ws2.append([seq, "中性", "/", "否", content, platform])
            c_empty += 1
    print(f"\n评论: 有标签 {c_lab} 条 / 空 {c_empty} 条", flush=True)

    # ── 重建 sheet3：一标签一行 ──
    ws3.delete_rows(2, ws3.max_row - 1)
    p_lab = p_empty = 0
    for r in post_rows:
        seq, title, platform = r[0], r[1], r[2]
        content = str(r[3]).strip() if r[3] else ""
        url = r[7] if len(r) > 7 else None
        expanded = post_cache.get(content, [])
        if expanded:
            for item in expanded:
                ws3.append([seq, title, platform, content, item["sentiment"], item["tag"], item["has_comparison"], url])
            p_lab += 1
        else:
            ws3.append([seq, title, platform, content, "/", "/", "否", url])
            p_empty += 1
    print(f"正文: 有标签 {p_lab} 条 / 空 {p_empty} 条", flush=True)

    wb.save(DST)
    print(f"\n已保存: {DST}", flush=True)


if __name__ == "__main__":
    main()
