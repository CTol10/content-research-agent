# -*- coding: utf-8 -*-
"""Compare classifier output with existing labels in the Excel file.

Filters rows by cell background color:
- 评论+标签+情感 sheet: blue (#5B9BD5) in B/C columns
- 正文内容+标签+情感 sheet: green (#70AD47) in E/F/G columns

Outputs comparison Excel with two sheets.
"""
import json
import sys
import time
from collections import defaultdict
from datetime import datetime

sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from openpyxl.styles import PatternFill

from classifier import classify_content, TAGS

EXCEL_PATH = r"input/【0529-0531】正文+评论.xlsx"
OUTPUT_DIR = "output"
BLUE_COLOR = "FF5B9BD5"
GREEN_COLOR = "FF70AD47"


def is_blue(cell):
    """Check if a cell has blue background fill."""
    if cell.fill and cell.fill.fill_type and cell.fill.fgColor:
        return cell.fill.fgColor.rgb == BLUE_COLOR
    return False


def is_green(cell):
    """Check if a cell has green background fill."""
    if cell.fill and cell.fill.fill_type and cell.fill.fgColor:
        return cell.fill.fgColor.rgb == GREEN_COLOR
    return False


def read_comment_sheet(wb):
    """Read 评论+标签+情感 sheet, filter rows with blue B/C cells."""
    ws = wb["评论+标签+情感"]
    rows = []
    content_seen = {}  # content -> first row data, to avoid duplicate API calls

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row_idx = row[0].row
        seq_cell, sent_cell, tag_cell, comp_cell, content_cell, platform_cell = row[0], row[1], row[2], row[3], row[4], row[5]

        # Only process rows where B or C has blue fill
        if not (is_blue(sent_cell) or is_blue(tag_cell)):
            continue

        content = content_cell.value
        if not content or not str(content).strip():
            continue

        content_text = str(content).strip()
        rows.append({
            "row_idx": row_idx,
            "seq": seq_cell.value,
            "orig_sentiment": str(sent_cell.value or "/").strip(),
            "orig_tag": str(tag_cell.value or "/").strip(),
            "orig_comparison": str(comp_cell.value or "否").strip(),
            "content": content_text,
            "platform": str(platform_cell.value or "").strip(),
        })

        # Track unique content for API dedup
        if content_text not in content_seen:
            content_seen[content_text] = len(rows) - 1

    return rows, content_seen


def read_post_sheet(wb):
    """Read 正文内容+标签+情感 sheet, filter rows with green E/F/G cells."""
    ws = wb["正文内容+标签+情感"]
    rows = []
    content_seen = {}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row_idx = row[0].row
        seq_cell = row[0]
        title_cell = row[1]
        platform_cell = row[2]
        content_cell = row[3]
        sent_cell = row[4]
        tag_cell = row[5]
        comp_cell = row[6]
        url_cell = row[7] if len(row) > 7 else None

        # Only process rows where E, F, or G has green fill
        if not (is_green(sent_cell) or is_green(tag_cell) or is_green(comp_cell)):
            continue

        content = content_cell.value
        if not content or not str(content).strip():
            continue

        content_text = str(content).strip()
        rows.append({
            "row_idx": row_idx,
            "seq": seq_cell.value,
            "title": str(title_cell.value or "").strip(),
            "platform": str(platform_cell.value or "").strip(),
            "content": content_text,
            "orig_sentiment": str(sent_cell.value or "/").strip(),
            "orig_tag": str(tag_cell.value or "/").strip(),
            "orig_comparison": str(comp_cell.value or "否").strip(),
            "url": str(url_cell.value or "").strip() if url_cell else "",
        })

        if content_text not in content_seen:
            content_seen[content_text] = len(rows) - 1

    return rows, content_seen


def clean_tag(raw_tag: str) -> str:
    """Extract valid tag name from potentially messy string."""
    raw_tag = raw_tag.strip()
    for tag in TAGS:
        if raw_tag.startswith(tag):
            return tag
    return raw_tag


def classify_and_parse(content_text):
    """Call MiMo API and return parsed result."""
    result = classify_content(content_text)
    tags = []
    sentiments = []
    for t in result["tags"]:
        parts = t.rsplit("-", 1)
        if len(parts) == 2:
            tags.append(parts[0])
            sentiments.append(parts[1])
    return {
        "tags": tags,
        "sentiments": sentiments,
        "has_comparison": result["has_comparison"],
    }


def compare_single(orig_tag, orig_sentiment, orig_comparison, new_result):
    """Compare one row's original values with new MiMo result."""
    # Normalize original tag
    orig_tag_clean = clean_tag(orig_tag) if orig_tag and orig_tag != "/" else ""
    orig_tags_set = {orig_tag_clean} if orig_tag_clean else set()

    new_tags_set = set(new_result["tags"])

    # Tag comparison
    if not orig_tags_set and "行业观察" in new_tags_set:
        tag_match = True
    else:
        tag_match = orig_tags_set == new_tags_set

    # Sentiment comparison
    orig_sent = orig_sentiment if orig_sentiment and orig_sentiment != "/" else ""
    new_sents = new_result["sentiments"]
    # Treat empty sentiment as 中性 (filtered comments are inherently neutral)
    if not new_sents:
        new_sents = ["中性"]
    # For single-tag results, compare directly
    if len(new_sents) == 1:
        sentiment_match = orig_sent == new_sents[0]
    elif orig_sent in new_sents:
        sentiment_match = True
    else:
        sentiment_match = False

    # Comparison match
    comp_match = orig_comparison == new_result["has_comparison"]

    missing_tags = sorted(orig_tags_set - new_tags_set)
    extra_tags = sorted(new_tags_set - orig_tags_set)

    return {
        "tag_match": tag_match,
        "sentiment_match": sentiment_match,
        "comp_match": comp_match,
        "missing_tags": missing_tags,
        "extra_tags": extra_tags,
    }


def write_output_excel(comment_results, post_results, output_path):
    """Write comparison results to Excel with two sheets."""
    wb = openpyxl.Workbook()

    # ── Sheet 1: 评论评价对比 ──
    ws1 = wb.active
    ws1.title = "评论评价对比"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    match_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    diff_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    headers1 = [
        "序号", "内容", "平台",
        "现有情感偏向", "新情感偏向", "情感一致",
        "现有标签", "新标签", "标签一致",
        "现有比较", "新比较", "比较一致",
    ]
    for col, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.fill = header_fill

    for i, r in enumerate(comment_results, 2):
        ws1.cell(row=i, column=1, value=r["seq"])
        ws1.cell(row=i, column=2, value=r["content"][:100])
        ws1.cell(row=i, column=3, value=r["platform"])
        ws1.cell(row=i, column=4, value=r["orig_sentiment"])
        ws1.cell(row=i, column=5, value=", ".join(r["new_sentiments"]) if r["new_sentiments"] else "/")
        ws1.cell(row=i, column=6, value="一致" if r["sentiment_match"] else "不一致")
        ws1.cell(row=i, column=7, value=r["orig_tag"])
        ws1.cell(row=i, column=8, value=", ".join(r["new_tags"]) if r["new_tags"] else "/")
        ws1.cell(row=i, column=9, value="一致" if r["tag_match"] else "不一致")
        ws1.cell(row=i, column=10, value=r["orig_comparison"])
        ws1.cell(row=i, column=11, value=r["new_comparison"])
        ws1.cell(row=i, column=12, value="一致" if r["comp_match"] else "不一致")

        # Color the match columns
        for col_idx in [6, 9, 12]:
            cell = ws1.cell(row=i, column=col_idx)
            cell.fill = match_fill if cell.value == "一致" else diff_fill

    # Auto-width
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws1.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    # ── Sheet 2: 正文评价对比 ──
    ws2 = wb.create_sheet("正文评价对比")

    headers2 = [
        "序号", "标题", "内容摘要",
        "现有情感偏向", "新情感偏向", "情感一致",
        "现有标签", "新标签", "标签一致",
        "现有比较", "新比较", "比较一致",
    ]
    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill

    for i, r in enumerate(post_results, 2):
        ws2.cell(row=i, column=1, value=r["seq"])
        ws2.cell(row=i, column=2, value=r["title"][:60])
        ws2.cell(row=i, column=3, value=r["content"][:100])
        ws2.cell(row=i, column=4, value=r["orig_sentiment"])
        ws2.cell(row=i, column=5, value=", ".join(r["new_sentiments"]) if r["new_sentiments"] else "/")
        ws2.cell(row=i, column=6, value="一致" if r["sentiment_match"] else "不一致")
        ws2.cell(row=i, column=7, value=r["orig_tag"])
        ws2.cell(row=i, column=8, value=", ".join(r["new_tags"]) if r["new_tags"] else "/")
        ws2.cell(row=i, column=9, value="一致" if r["tag_match"] else "不一致")
        ws2.cell(row=i, column=10, value=r["orig_comparison"])
        ws2.cell(row=i, column=11, value=r["new_comparison"])
        ws2.cell(row=i, column=12, value="一致" if r["comp_match"] else "不一致")

        for col_idx in [6, 9, 12]:
            cell = ws2.cell(row=i, column=col_idx)
            cell.fill = match_fill if cell.value == "一致" else diff_fill

    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws2.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    wb.save(output_path)
    print(f"\n结果已保存到: {output_path}")


def run_comparison():
    print(f"读取文件: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # ── Read and filter ──
    print("读取评论 sheet (蓝色底色行)...")
    comment_rows, comment_content_map = read_comment_sheet(wb)
    print(f"  筛选出 {len(comment_rows)} 行 (去重后 {len(comment_content_map)} 条内容)")

    print("读取正文 sheet (绿色底色行)...")
    post_rows, post_content_map = read_post_sheet(wb)
    print(f"  筛选出 {len(post_rows)} 行 (去重后 {len(post_content_map)} 条内容)")

    wb.close()

    # ── Classify comments ──
    print(f"\n{'='*60}")
    print(f"评论分类 ({len(comment_content_map)} 条内容需要 API 调用)")
    print(f"{'='*60}")

    comment_cache = {}
    for i, (content, idx) in enumerate(comment_content_map.items()):
        print(f"  [{i+1}/{len(comment_content_map)}] {content[:40]}...", end=" ", flush=True)
        try:
            result = classify_and_parse(content)
            comment_cache[content] = result
            print(f"标签={result['tags']} 情感={result['sentiments']} 对比={result['has_comparison']}")
        except Exception as e:
            print(f"ERROR: {e}")
            comment_cache[content] = {"tags": [], "sentiments": [], "has_comparison": "否"}
        time.sleep(0.3)

    # Build comment results
    comment_results = []
    s2_sent_match = s2_tag_match = s2_comp_match = 0
    for r in comment_rows:
        new_result = comment_cache.get(r["content"], {"tags": [], "sentiments": [], "has_comparison": "否"})
        cmp = compare_single(r["orig_tag"], r["orig_sentiment"], r["orig_comparison"], new_result)

        entry = {
            **r,
            "new_tags": new_result["tags"],
            "new_sentiments": new_result["sentiments"],
            "new_comparison": new_result["has_comparison"],
            **cmp,
        }
        comment_results.append(entry)
        if cmp["sentiment_match"]:
            s2_sent_match += 1
        if cmp["tag_match"]:
            s2_tag_match += 1
        if cmp["comp_match"]:
            s2_comp_match += 1

    # ── Classify posts ──
    print(f"\n{'='*60}")
    print(f"正文分类 ({len(post_content_map)} 条内容需要 API 调用)")
    print(f"{'='*60}")

    post_cache = {}
    for i, (content, idx) in enumerate(post_content_map.items()):
        title = post_rows[idx]["title"]
        print(f"  [{i+1}/{len(post_content_map)}] {title[:40]}...", end=" ", flush=True)
        try:
            result = classify_and_parse(content)
            post_cache[content] = result
            print(f"标签={result['tags']} 情感={result['sentiments']} 对比={result['has_comparison']}")
        except Exception as e:
            print(f"ERROR: {e}")
            post_cache[content] = {"tags": [], "sentiments": [], "has_comparison": "否"}
        time.sleep(0.3)

    # Build post results
    post_results = []
    s3_sent_match = s3_tag_match = s3_comp_match = 0
    for r in post_rows:
        new_result = post_cache.get(r["content"], {"tags": [], "sentiments": [], "has_comparison": "否"})
        cmp = compare_single(r["orig_tag"], r["orig_sentiment"], r["orig_comparison"], new_result)

        entry = {
            **r,
            "new_tags": new_result["tags"],
            "new_sentiments": new_result["sentiments"],
            "new_comparison": new_result["has_comparison"],
            **cmp,
        }
        post_results.append(entry)
        if cmp["sentiment_match"]:
            s3_sent_match += 1
        if cmp["tag_match"]:
            s3_tag_match += 1
        if cmp["comp_match"]:
            s3_comp_match += 1

    # ── Summary ──
    total_comment = len(comment_results)
    total_post = len(post_results)

    print(f"\n{'='*60}")
    print("汇总")
    print(f"{'='*60}")
    print(f"评论 ({total_comment} 行):")
    print(f"  情感一致: {s2_sent_match}/{total_comment} ({s2_sent_match/total_comment*100:.1f}%)" if total_comment else "  无数据")
    print(f"  标签一致: {s2_tag_match}/{total_comment} ({s2_tag_match/total_comment*100:.1f}%)" if total_comment else "")
    print(f"  比较一致: {s2_comp_match}/{total_comment} ({s2_comp_match/total_comment*100:.1f}%)" if total_comment else "")
    print(f"正文 ({total_post} 行):")
    print(f"  情感一致: {s3_sent_match}/{total_post} ({s3_sent_match/total_post*100:.1f}%)" if total_post else "  无数据")
    print(f"  标签一致: {s3_tag_match}/{total_post} ({s3_tag_match/total_post*100:.1f}%)" if total_post else "")
    print(f"  比较一致: {s3_comp_match}/{total_post} ({s3_comp_match/total_post*100:.1f}%)" if total_post else "")

    # ── Write output ──
    import os
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    date_str = datetime.now().strftime("%Y%m%d")
    output_path = os.path.join(OUTPUT_DIR, f"label_comparison_{date_str}.xlsx")
    write_output_excel(comment_results, post_results, output_path)

    # Also save JSON for reference
    json_path = os.path.join(OUTPUT_DIR, f"label_comparison_{date_str}.json")
    json_data = {
        "summary": {
            "comments": {"total": total_comment, "sentiment_match": s2_sent_match, "tag_match": s2_tag_match, "comp_match": s2_comp_match},
            "posts": {"total": total_post, "sentiment_match": s3_sent_match, "tag_match": s3_tag_match, "comp_match": s3_comp_match},
        },
        "comment_results": comment_results,
        "post_results": post_results,
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(json_data, f, ensure_ascii=False, indent=2, default=str)
    print(f"JSON 详情已保存到: {json_path}")


if __name__ == "__main__":
    run_comparison()
