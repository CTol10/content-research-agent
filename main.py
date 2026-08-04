import argparse
import asyncio
import importlib
import logging
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from classifier import classify_content_expanded, TAGS
from scrapers.comment_cleaner import clean_comment_content

logger = logging.getLogger("scraper")

# ── Dependency check & auto-install ──────────────────────────────────────

REQUIRED_PACKAGES = {
    "playwright": "playwright",
    "openpyxl": "openpyxl",
}


def check_dependencies() -> list[str]:
    """Return list of missing package names."""
    missing = []
    for module_name, pip_name in REQUIRED_PACKAGES.items():
        try:
            importlib.import_module(module_name)
        except ImportError:
            missing.append(pip_name)
    return missing


def _detect_browser_channel() -> str | None:
    """Detect available system browser. Returns channel name or None."""
    # In frozen exe, skip Playwright's own browser check (it can't find its bundled Chromium)
    # and detect system Chrome/Edge directly
    if getattr(sys, 'frozen', False):
        import winreg
        for reg_key, channel in [
            (r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\chrome.exe", "chrome"),
            (r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\msedge.exe", "msedge"),
        ]:
            try:
                key = winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, reg_key)
                val, _ = winreg.QueryValueEx(key, "")
                winreg.CloseKey(key)
                if val and Path(val).exists():
                    return channel
            except (OSError, FileNotFoundError):
                pass
        # Fallback: check common install paths
        for path_pattern, channel in [
            (r"C:\Program Files\Google\Chrome\Application\chrome.exe", "chrome"),
            (r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe", "msedge"),
        ]:
            if Path(path_pattern).exists():
                return channel
        return None

    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            try:
                b = p.chromium.launch(channel="chrome", headless=True)
                b.close()
                return "chrome"
            except Exception:
                pass
            try:
                b = p.chromium.launch(channel="msedge", headless=True)
                b.close()
                return "msedge"
            except Exception:
                pass
    except Exception:
        pass
    return None


def check_playwright_browser() -> bool:
    """Check if system browser or Chromium is available."""
    if _detect_browser_channel():
        return True
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            browser.close()
        return True
    except Exception:
        return False


def _install_chromium_with_fallback() -> bool:
    """Install Chromium browser. Tries default CDN first, then npmmirror."""
    try:
        subprocess.check_call(
            [sys.executable, "-m", "playwright", "install", "chromium"],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        return True
    except subprocess.CalledProcessError:
        pass

    mirror = "https://cdn.npmmirror.com/binaries/playwright"
    print(f"  -> 默认 CDN 安装失败，尝试国内镜像: {mirror}")
    env = os.environ.copy()
    env["PLAYWRIGHT_DOWNLOAD_HOST"] = mirror
    try:
        subprocess.check_call(
            [sys.executable, "-m", "playwright", "install", "chromium"],
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        return True
    except subprocess.CalledProcessError:
        pass

    print("  -> 国内镜像版本可能不匹配，请手动安装:")
    print("     1. 使用 VPN/代理后运行: playwright install chromium")
    print("     2. 或手动下载: https://playwright.dev/docs/browsers")
    return False


def install_dependencies():
    """Install missing pip packages and Playwright browser."""
    if getattr(sys, 'frozen', False):
        print("\n[提示] 当前为打包版本，跳过依赖安装。")
        return True
    print("\n" + "=" * 60)
    print("  环境检查与自动安装")
    print("=" * 60)

    missing = check_dependencies()
    if missing:
        print(f"\n[1/3] 正在安装缺失的 Python 包: {', '.join(missing)}")
        try:
            subprocess.check_call(
                [sys.executable, "-m", "pip", "install", "--quiet"] + missing,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )
            print(f"  -> 安装成功: {', '.join(missing)}")
        except subprocess.CalledProcessError as e:
            print(f"  -> 安装失败: {e}")
            print(f"  -> 请手动运行: pip install {' '.join(missing)}")
            return False
    else:
        print("\n[1/3] Python 包已就绪")

    print("[2/3] 检查浏览器...")
    channel = _detect_browser_channel()
    if channel:
        print(f"  -> 检测到系统 {channel}，无需下载 Chromium")
    elif not check_playwright_browser():
        print("  -> 未检测到系统 Chrome，正在安装 Chromium 浏览器（约 300MB）...")
        if _install_chromium_with_fallback():
            print("  -> Chromium 安装成功")
        else:
            print("  -> Chromium 安装失败")
            print("  -> 请确保已安装 Google Chrome 或手动运行: playwright install chromium")
            return False
    else:
        print("  -> Chromium 已就绪")

    print("\n" + "=" * 60)
    print("  环境检查通过！可以开始使用了。")
    print("=" * 60)
    return True


# ── Logging ──────────────────────────────────────────────────────────────

def setup_logging():
    import config
    config.LOG_DIR.mkdir(parents=True, exist_ok=True)
    log_file = config.LOG_DIR / f"scrape_{datetime.now():%Y%m%d}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[
            logging.FileHandler(log_file, encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )


# ── Excel I/O ───────────────────────────────────────────────────────────

def read_input_excel(filepath) -> list[dict]:
    """Read input Excel and return rows with valid URLs and supported platforms."""
    import config
    from scrapers import match_platform

    if not Path(filepath).exists():
        print(f"\n[错误] 找不到输入文件: {filepath}")
        print("请检查文件路径是否正确，或者使用 --input 指定文件路径")
        return []

    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True)
    except Exception as e:
        print(f"\n[错误] 无法打开 Excel 文件: {e}")
        print("请确认文件是 .xlsx 格式，且没有被其他程序占用")
        return []

    # Prefer "原始底表" sheet if it exists (matches expected column layout)
    if "原始底表" in wb.sheetnames:
        ws = wb["原始底表"]
    else:
        ws = wb.active
    rows = []
    skipped_urls = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        url = row[config.COL_URL]
        if not url or not isinstance(url, str) or not url.startswith("http"):
            continue
        platform_key = match_platform(url)
        if platform_key is None:
            domain = str(url).split('/')[2] if len(str(url).split('/')) > 2 else str(url)[:60]
            skipped_urls.append(domain)
            continue
        rows.append({
            "serial": row[config.COL_SERIAL],
            "title": str(row[config.COL_TITLE] or ""),
            "source": str(row[config.COL_SOURCE] or ""),
            "author": str(row[config.COL_AUTHOR] or ""),
            "url": url,
            "platform_key": platform_key,
        })
    wb.close()

    if skipped_urls:
        from collections import Counter
        url_counts = Counter(skipped_urls)
        print(f"\n[提示] 跳过 {len(skipped_urls)} 条不支持的链接（域名不在 PLATFORM_PATTERNS 中）：")
        for domain, count in url_counts.most_common():
            print(f"  {domain}: {count} 条")

    if not rows:
        print("\n[警告] Excel 中没有找到支持的链接")
        print("支持的平台: 抖音(douyin), 小红书(xiaohongshu), 微博(weibo), 今日头条(toutiao)")
        print("请确认 Excel 第3列(C列)包含有效链接")

    return rows


def _resolve_config_path(filename: str) -> Path | None:
    """Resolve a config file path — checks exe directory first, then CWD.

    scraper 实际按 config.BASE_DIR（exe 同目录）读取；这里保持同序，
    避免 CWD≠exe 目录时误判配置缺失（例如从其他目录调起 exe）。
    """
    import config
    base_path = config.BASE_DIR / filename
    if base_path.exists():
        return base_path
    cwd_path = Path(filename)
    if cwd_path.exists():
        return cwd_path
    return None


def _ensure_wechat_config(requested_platforms: set):
    """Auto-check WeChat PC calibration config. Prompt to calibrate if missing."""
    wechat_keys = {"wechat", "wechat_channels"}
    # Only check when wechat platforms are explicitly requested, or when
    # running without platform filter (all platforms — check anyway)
    if requested_platforms and not (requested_platforms & wechat_keys):
        return

    from scrapers.wechat import detect_version
    version = detect_version()

    if version == "4.1.7":
        config_filename = "config.wechat_pc_417.json"
        alt_config_filename = "config.wechat_pc.json"
        calibrator_module = "scrapers.wechat.v417.calibrator"
        calibrator_cls = "WechatCalibratorV417"
        alt_version = "4.1.11"
    else:
        config_filename = "config.wechat_pc.json"
        alt_config_filename = "config.wechat_pc_417.json"
        calibrator_module = "scrapers.wechat_calibrator"
        calibrator_cls = "WechatCalibrator"
        alt_version = "4.1.7"

    config_path = _resolve_config_path(config_filename)

    # If the detected version's config doesn't exist, try the other version's config
    if config_path is None:
        config_path = _resolve_config_path(alt_config_filename)
        if config_path is not None:
            logger.info("检测到 %s 配置文件，切换为 %s 配置", alt_config_filename, alt_version)
            if alt_version == "4.1.7":
                calibrator_module = "scrapers.wechat.v417.calibrator"
                calibrator_cls = "WechatCalibratorV417"
            else:
                calibrator_module = "scrapers.wechat_calibrator"
                calibrator_cls = "WechatCalibrator"
            version = alt_version

    if config_path is None:
        print("\n" + "=" * 60)
        print(f"  检测到需要微信公众号/视频号功能（WeChat {version}），但未找到坐标配置文件。")
        print("  需要先校准微信窗口中的 UI 元素位置。")
        print("=" * 60)
        _offer_calibration_v2(version, calibrator_module, calibrator_cls)
        return

    try:
        import json
        data = json.loads(config_path.read_text(encoding="utf-8"))
        profile_name = data.get("active_profile", "")
        profile = data.get("profiles", {}).get(profile_name, {})
        if not profile or "wechat_main" not in profile:
            print("\n" + "=" * 60)
            print(f"  微信坐标配置不完整（WeChat {version}），需要重新校准。")
            print("=" * 60)
            _offer_calibration_v2(version, calibrator_module, calibrator_cls)
    except Exception:
        print("\n" + "=" * 60)
        print(f"  微信坐标配置文件损坏（WeChat {version}），需要重新校准。")
        print("=" * 60)
        _offer_calibration_v2(version, calibrator_module, calibrator_cls)


def _offer_calibration_v2(version: str, module_name: str, class_name: str):
    """Offer to run version-appropriate WeChat calibration."""
    print(f"  检测到 WeChat {version}，将使用对应版本的校准工具。")
    print("  请确保微信 PC 客户端已启动并登录。")
    print()
    try:
        answer = input("  是否现在开始校准？[Y/n] ").strip().lower()
    except (EOFError, KeyboardInterrupt):
        print("\n  已取消。请稍后手动运行校准工具。")
        sys.exit(1)
    if answer and answer not in ("y", "yes", ""):
        print("  已跳过校准。请稍后手动运行校准工具。")
        sys.exit(1)
    print()
    import importlib
    mod = importlib.import_module(module_name)
    calibrator_cls = getattr(mod, class_name)
    calibrator = calibrator_cls()
    ok = calibrator.calibrate_all()
    if not ok:
        print("\n  校准未完成，无法继续。")
        sys.exit(1)
    print("\n  校准完成，继续抓取流程...\n")


def run_wechat_calibration(use_overlay: bool = True) -> bool:
    """Run version-appropriate WeChat PC coordinate calibration."""
    from scrapers.wechat import detect_version
    version = detect_version()
    if version == "4.1.7":
        from scrapers.wechat.v417.calibrator import WechatCalibratorV417
        calibrator = WechatCalibratorV417(use_overlay=use_overlay)
    else:
        from scrapers.wechat_calibrator import WechatCalibrator
        calibrator = WechatCalibrator(use_overlay=use_overlay)
    return bool(calibrator.calibrate_all())


def _show_startup_menu(args) -> bool:
    """裸运行（无动作参数）启动菜单：抓取 / 校准 / 退出。

    Returns True 表示继续抓取，False 表示退出。
    """
    while True:
        print()
        print("=" * 60)
        print("  请选择操作：")
        print("    [1] 开始抓取")
        print("    [2] 校准微信坐标（内置默认坐标不匹配时使用）")
        print("    [3] 退出")
        print("=" * 60)
        try:
            choice = input("  请选择 [1/2/3]（默认 1）: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\n  已退出。")
            return False
        if choice in ("", "1"):
            return True
        if choice == "2":
            ok = run_wechat_calibration(use_overlay=not args.no_overlay)
            if ok:
                print("\n  校准完成。可重新选择开始抓取。\n")
            else:
                print("\n  校准未完成，可重试或退出。\n")
            continue
        if choice == "3":
            print("  再见！")
            return False
        print("  无效输入，请重新选择。")


_PLATFORM_DISPLAY = {
    "douyin": "抖音",
    "xiaohongshu": "小红书",
    "weibo": "微博",
    "toutiao": "今日头条",
    "wechat": "微信公众号",
    "wechat_channels": "微信视频号",
    "bilibili": "哔哩哔哩",
}


def init_output_workbook(input_file: str):
    import openpyxl
    wb = openpyxl.Workbook()
    src_wb = openpyxl.load_workbook(input_file)
    src_ws = src_wb["原始底表"] if "原始底表" in src_wb.sheetnames else src_wb.active
    ws1 = wb.active
    ws1.title = "原始底表"
    for row in src_ws.iter_rows(values_only=True):
        ws1.append(list(row))
    src_wb.close()

    ws2 = wb.create_sheet("评论+标签+情感")
    ws2.append(["序号", "情感偏向", "标签", "是否与其他酒店比较", "内容", "平台"])

    ws3 = wb.create_sheet("正文内容+标签+情感")
    ws3.append(["序号", "标题", "平台", "内容摘要", "情感偏向", "标签", "是否与其他酒店比较", "原文/评论链接"])
    return wb


def write_summary_report(output_dir: Path, results: list[dict]):
    import openpyxl
    filepath = output_dir / f"scrape_report_{datetime.now():%Y%m%d}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["序号", "平台", "标题", "URL", "状态", "评论数", "错误信息"])
    for idx, r in enumerate(results, start=1):
        ws.append([
            idx,
            r.get("platform", ""),
            r.get("title", ""),
            r.get("url", ""),
            r.get("status", ""),
            r.get("comment_count", 0),
            r.get("error", ""),
        ])
    wb.save(filepath)
    logger.info(f"汇总报告已保存: {filepath}")


# ── Login ────────────────────────────────────────────────────────────────

CDP_PLATFORMS = set()


def check_cookies(platform_keys: list[str]) -> list[str]:
    import config
    # WeChat platforms use desktop automation, not browser cookies.
    # A separate WeChat window check runs before scraping.
    _WECHAT_PLATFORMS = {"wechat", "wechat_channels"}
    missing = []
    for key in platform_keys:
        if key in _WECHAT_PLATFORMS:
            continue
        # Check if cookie file exists OR persistent browser profile exists
        cookie_path = config.COOKIE_DIR / f"{key}.json"
        browser_profile = config.COOKIE_DIR / "_browser_profile"
        if not cookie_path.exists() and not browser_profile.exists():
            missing.append(key)
    return missing


async def login_platform(platform_key: str, wait_seconds: int = 300):
    from scrapers.douyin import DouyinScraper
    from scrapers.xiaohongshu import XiaohongshuScraper
    from scrapers.weibo import WeiboScraper
    from scrapers.toutiao import ToutiaoScraper
    from scrapers.bilibili import BilibiliScraper

    scraper_map = {
        "douyin": DouyinScraper,
        "xiaohongshu": XiaohongshuScraper,
        "weibo": WeiboScraper,
        "toutiao": ToutiaoScraper,
        "bilibili": BilibiliScraper,
    }
    scraper_cls = scraper_map.get(platform_key)
    if not scraper_cls:
        print(f"[错误] 未知平台: {platform_key}，支持的平台: {', '.join(scraper_map.keys())}")
        return
    scraper = scraper_cls()
    await scraper.wait_for_login(timeout=wait_seconds)


async def login_open_all(platform_keys: list[str]):
    """Open browsers for all missing platforms, poll for signal files to save cookies."""
    from scrapers.douyin import DouyinScraper
    from scrapers.xiaohongshu import XiaohongshuScraper
    from scrapers.weibo import WeiboScraper
    from scrapers.toutiao import ToutiaoScraper
    from scrapers.bilibili import BilibiliScraper

    scraper_map = {
        "douyin": DouyinScraper,
        "xiaohongshu": XiaohongshuScraper,
        "weibo": WeiboScraper,
        "toutiao": ToutiaoScraper,
        "bilibili": BilibiliScraper,
    }

    missing = check_cookies(platform_keys)
    if not missing:
        print("\n所有平台均已登录，无需重复登录。")
        return

    print(f"\n需要登录的平台: {', '.join(missing)}")
    print("将依次打开浏览器，请在浏览器中手动登录。\n")

    for key in missing:
        scraper_cls = scraper_map.get(key)
        if not scraper_cls:
            continue
        try:
            scraper = scraper_cls()
            await scraper.login_interactive(wait_seconds=0)
        except Exception as e:
            print(f"[{key}] 打开失败: {e}")

    print("\n所有平台登录完成。")


async def login_save_all(platform_keys: list[str]):
    """Write signal files to trigger cookie saves in running login processes."""
    missing = check_cookies(platform_keys)
    if not missing:
        print("\n所有平台均已登录，无需重复保存。")
        return

    print(f"\n正在保存以下平台的 Cookie: {', '.join(missing)}\n")

    for key in missing:
        signal_file = config.COOKIE_DIR / f".{key}_login_done"
        signal_file.write_text("done", encoding="utf-8")
        print(f"[{key}] 已发送保存信号")

    print("\n所有 Cookie 保存完成！可以开始抓取了。")


# ── Scrape ───────────────────────────────────────────────────────────────

ERROR_MESSAGES = {
    "no_comments": "该视频暂无评论",
    "note_unavailable": "笔记不可用（可能已被删除或设为私密）",
    "note_anti_crawl": "笔记触发反爬机制（已自动重试）",
    "login_required": "需要重新登录（Cookie 已过期）",
    "network_error": "网络连接失败，请检查网络后重试",
    "page_timeout": "页面加载超时，可能网络较慢或平台繁忙",
    "unknown": "未知错误",
}


def classify_error(error_str: str) -> str:
    """Map raw error to user-friendly Chinese message."""
    lower = error_str.lower()
    if "暂无评论" in error_str:
        return ERROR_MESSAGES["no_comments"]
    if "笔记不可用" in error_str or "noteunavailable" in lower:
        return ERROR_MESSAGES["note_unavailable"]
    if "反爬" in error_str:
        return ERROR_MESSAGES["note_anti_crawl"]
    if "login" in lower or "cookie" in lower:
        return ERROR_MESSAGES["login_required"]
    if "timeout" in lower or "timed out" in lower:
        return ERROR_MESSAGES["page_timeout"]
    if "connection" in lower or "connect" in lower:
        return ERROR_MESSAGES["network_error"]
    return error_str


async def run_all(input_file, platforms_filter=None, resume=False, batch_size=0,
                  login_timeout=300, classify=True, process_video=True, headless=False):
    """Full pipeline: check cookies → auto-login if needed → scrape."""
    import config
    from scrapers.douyin import DouyinScraper
    from scrapers.xiaohongshu import XiaohongshuScraper
    from scrapers.wechat import get_official_scraper, get_channels_scraper
    from scrapers.weibo import WeiboScraper
    from scrapers.toutiao import ToutiaoScraper
    from scrapers.bilibili import BilibiliScraper

    scraper_classes = {
        "douyin": DouyinScraper,
        "xiaohongshu": XiaohongshuScraper,
        "wechat": lambda: get_official_scraper(),
        "wechat_channels": lambda: get_channels_scraper(),
        "weibo": WeiboScraper,
        "toutiao": ToutiaoScraper,
        "bilibili": BilibiliScraper,
    }

    rows = read_input_excel(input_file)
    if not rows:
        print("[错误] Excel 中没有找到支持的链接")
        return

    if platforms_filter:
        rows = [r for r in rows if r["platform_key"] in platforms_filter]

    if not rows:
        print("[错误] 筛选后没有匹配的链接")
        return

    needed_platforms = list(set(r["platform_key"] for r in rows))
    missing = check_cookies(needed_platforms)

    if missing:
        print(f"\n{'='*50}")
        print(f"  需要登录的平台: {', '.join(missing)}")
        print(f"  浏览器将自动打开，请在浏览器中登录")
        print(f"  登录成功后会自动检测并保存 Cookie")
        print(f"{'='*50}\n")

        for key in missing:
            scraper_cls = scraper_classes.get(key)
            if not scraper_cls:
                continue
            scraper = scraper_cls()
            try:
                success = await scraper.wait_for_login(timeout=login_timeout)
                if not success:
                    print(f"[{key}] 登录超时，跳过该平台")
            except Exception as e:
                print(f"[{key}] 登录失败: {e}")

        still_missing = check_cookies(needed_platforms)
        if still_missing:
            print(f"\n[警告] 以下平台仍未登录: {', '.join(still_missing)}")

    print(f"\n开始抓取...\n")
    return await scrape_all(input_file, platforms_filter, resume=resume,
                            batch_size=batch_size, classify=classify,
                            process_video=process_video, headless=headless)


async def scrape_all(input_file, platforms_filter=None, resume=False, batch_size=0,
                     classify=True, process_video=True, headless=False):
    import json as _json
    import config
    from scrapers.douyin import DouyinScraper
    from scrapers.xiaohongshu import XiaohongshuScraper
    from scrapers.wechat import get_official_scraper, get_channels_scraper
    from scrapers.toutiao import ToutiaoScraper
    from scrapers.weibo import WeiboScraper
    from scrapers.bilibili import BilibiliScraper

    scraper_classes = {
        "douyin": DouyinScraper,
        "xiaohongshu": XiaohongshuScraper,
        "wechat": lambda: get_official_scraper(),
        "wechat_channels": lambda: get_channels_scraper(),
        "toutiao": ToutiaoScraper,
        "weibo": WeiboScraper,
        "bilibili": BilibiliScraper,
    }

    rows = read_input_excel(input_file)
    if not rows:
        return

    if classify and not config.MIMO_API_KEY:
        print("\n[警告] 未设置 MIMO_API_KEY 环境变量，分类功能将被跳过。")
        print("设置方法: set MIMO_API_KEY=your_api_key (Windows)")
        classify = False

    total = len(rows)
    logger.info(f"开始抓取，共 {total} 条链接")

    # ── Checkpoint loading ──────────────────────────────────────
    checkpoint_path = config.OUTPUT_DIR / "scrape_checkpoint.jsonl"
    completed_urls: set[str] = set()

    if resume and checkpoint_path.exists():
        try:
            lines = checkpoint_path.read_text(encoding="utf-8").splitlines()
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                item = _json.loads(line)
                completed_urls.add(item["url"])
            logger.info(f"从 checkpoint 恢复，已完成 {len(completed_urls)} 条")
            print(f"\n[续抓] 从 checkpoint 恢复，已完成 {len(completed_urls)} 条，跳过已抓链接")
        except Exception as e:
            logger.warning(f"读取 checkpoint 失败，从头开始: {e}")
            completed_urls = set()

    if platforms_filter:
        rows = [r for r in rows if r["platform_key"] in platforms_filter]
        logger.info(f"筛选平台 {platforms_filter} 后，剩余 {len(rows)} 条链接")
        if not rows:
            print(f"\n[警告] 筛选后没有匹配的链接")
            print(f"支持的平台: {', '.join(scraper_classes.keys())}")
            return

    # ── Filter out already-scraped URLs ─────────────────────────
    if completed_urls:
        before = len(rows)
        rows = [r for r in rows if r["url"] not in completed_urls]
        logger.info(f"跳过已完成 {before - len(rows)} 条，剩余 {len(rows)} 条")
        print(f"  跳过已完成 {before - len(rows)} 条，剩余 {len(rows)} 条待抓取")

    # ── Batch size limit ──────────────────────────────────────
    if batch_size > 0 and len(rows) > batch_size:
        rows = rows[:batch_size]
        logger.info(f"批次限制，本轮抓取 {batch_size} 条")
        print(f"  批次限制：本轮抓取前 {batch_size} 条")

    if not rows:
        print("\n所有链接均已抓取完成。")
        if checkpoint_path.exists():
            checkpoint_path.unlink()
        return

    needed_platforms = list(set(r["platform_key"] for r in rows))
    missing = check_cookies(needed_platforms)
    if missing:
        print(f"\n{'='*50}")
        print(f"  以下平台需要登录: {', '.join(missing)}")
        print(f"  浏览器将依次打开，请在每个浏览器窗口中完成登录")
        print(f"  登录成功后工具会自动检测并继续")
        print(f"{'='*50}\n")

        for idx, key in enumerate(missing, 1):
            scraper_cls = scraper_classes.get(key)
            if not scraper_cls:
                continue
            platform_display = _PLATFORM_DISPLAY.get(key, key)
            print(f"\n[{idx}/{len(missing)}] 正在打开 {platform_display} 登录页面...")
            scraper = scraper_cls()
            try:
                success = await scraper.wait_for_login(timeout=300)
                if not success:
                    print(f"[{platform_display}] 登录超时，该平台链接将被跳过")
            except Exception as e:
                print(f"[{platform_display}] 登录失败: {e}")

        # Re-check after login attempt
        still_missing = check_cookies(needed_platforms)
        if still_missing:
            print(f"\n[警告] 以下平台仍未登录: {', '.join(still_missing)}")
            print("这些平台的链接将被跳过。")

    # ── WeChat window check ──────────────────────────────────────
    wechat_needed = [p for p in needed_platforms if p in ("wechat", "wechat_channels")]
    if wechat_needed:
        if headless:
            print("\n" + "=" * 60)
            print("  [提示] 无头模式下不支持微信平台（需要 PC 客户端窗口）。")
            print(f"  将跳过微信相关链接: {', '.join(wechat_needed)}")
            print("=" * 60)
            rows = [r for r in rows if r["platform_key"] not in ("wechat", "wechat_channels")]
            if not rows:
                print("没有可抓取的非微信链接，退出。")
                return
        else:
            from scrapers.wechat import detect_version
            version = detect_version()
            if version == "4.1.7":
                from scrapers.wechat.v417.window_manager import WechatWindowManagerV417
                wm = WechatWindowManagerV417()
                window_found = wm.find_main()
            else:
                from scrapers.wechat_window_manager import WechatWindowManager
                wm = WechatWindowManager()
                window_found = wm.find_window()
            if not window_found:
                print("\n" + "=" * 60)
                print("  [提示] 检测到需要抓取微信平台，但微信窗口未打开。")
                print("  请打开微信 PC 客户端并登录，然后按 Enter 继续...")
                print("  （如果不需要抓取微信，可以按 Ctrl+C 退出，")
                print("   下次运行时使用 --platforms 参数排除微信平台）")
                print("=" * 60)
                try:
                    input()
                except (EOFError, KeyboardInterrupt):
                    print("\n用户取消。")
                    return

            # ── 元宝登录态检查（视频号口播在线解析所需）──
            # 直接跑抓取（--run）时若缺 _yuanbao_profile，提示扫码登录，不必单独 --login。
            if "wechat_channels" in needed_platforms:
                from scrapers.wechat.channels_sph_parser import ChannelsSphParser
                yuanbao_profile = config.COOKIE_DIR / "_yuanbao_profile"
                if not yuanbao_profile.exists():
                    print("\n" + "=" * 56)
                    print("  视频号口播在线解析需要腾讯元宝登录态")
                    print("  未检测到 cookies/_yuanbao_profile，现在扫码登录？[Y/n]")
                    print("=" * 56)
                    try:
                        _yb_ans = input().strip().lower()
                    except (EOFError, KeyboardInterrupt):
                        _yb_ans = ""
                    if _yb_ans in ("", "y", "yes"):
                        await ChannelsSphParser.login_interactive(wait_seconds=300)
                    else:
                        print("  跳过——视频号口播将为空（之后可用 --login 单独登录元宝）。")

    results = []
    success_count = 0
    fail_count = 0

    now = datetime.now()
    output_filepath = config.OUTPUT_DIR / f"scrape_result_{now:%Y%m%d_%H%M%S}.xlsx"
    wb = init_output_workbook(input_file)
    ws_comments = wb["评论+标签+情感"]
    ws_analysis = wb["正文内容+标签+情感"]
    analysis_counter = 1
    comment_counter = 1

    # Video processor initialization
    video_processor = None
    if process_video:
        from video_processor import VideoProcessor
        if not config.MIMO_API_KEY:
            print("\n[警告] 未设置 MIMO_API_KEY，视频处理将被跳过。")
            process_video = False
        else:
            video_processor = VideoProcessor()

    # Collect video URLs for inspection
    video_urls_log = []

    active_platforms = {r["platform_key"] for r in rows}

    # Ensure WeChat calibration is ready if wechat platforms are active
    _ensure_wechat_config(active_platforms)

    for platform_key, scraper_cls in scraper_classes.items():
        if platform_key not in active_platforms:
            continue

        platform_rows = [r for r in rows if r["platform_key"] == platform_key]
        platform_display = _PLATFORM_DISPLAY.get(platform_key, platform_key)
        print(f"\n--- 开始抓取 {platform_display} ({len(platform_rows)} 条链接) ---")

        scraper = scraper_cls()
        try:
            is_wechat_p = platform_key in ("wechat", "wechat_channels")
            start_kwargs = {} if is_wechat_p else {"headless": headless}
            await scraper.start(**start_kwargs)
        except Exception as e:
            is_wechat = platform_key in ("wechat", "wechat_channels")
            if is_wechat:
                print(f"[错误] 启动微信抓取失败: {e}")
                print("请确保微信 PC 客户端已启动并登录，然后重试。")
            else:
                print(f"[错误] 无法启动浏览器: {e}")
                print("请确认已安装 Playwright Chromium: playwright install chromium")
            fail_count += len(platform_rows)
            for row in platform_rows:
                results.append({
                    "platform": platform_key,
                    "title": row["title"],
                    "url": row["url"],
                    "status": "失败",
                    "comment_count": 0,
                    "error": f"微信启动失败: {e}" if is_wechat else f"浏览器启动失败: {e}",
                })
            continue

        try:
            for i, row in enumerate(platform_rows, 1):
                url = row["url"]
                title_display = row['title'][:30].encode('gbk', errors='replace').decode('gbk')
                print(f"  [{i}/{len(platform_rows)}] {title_display}...", end=" ", flush=True)

                try:
                    result = await scraper.scrape(url)
                    post_content = result.get("post_content", "")
                    comments = result.get("comments", [])
                    count = len(comments)

                    # Log video URL
                    video_url = result.get("video_url")
                    if video_url:
                        video_urls_log.append({"url": url, "video_url": video_url, "platform": platform_key})
                        print(f"    [视频] {video_url[:80]}")

                    platform_name = _PLATFORM_DISPLAY.get(platform_key, platform_key)

                    # Write comments to Excel (one row per tag)
                    written_in_url = 0
                    dropped_in_url = 0
                    for ci, (_, content) in enumerate(comments):
                        # Clean comment content: strip reply prefix, metadata, etc.
                        raw_repr = repr(content)[:80] if content is not None else "None"
                        cleaned_content = clean_comment_content(content, platform=platform_key)
                        if not cleaned_content:
                            dropped_in_url += 1
                            logger.warning(
                                f"[{platform_key}] 评论被清洗丢弃 ci={ci+1}/{len(comments)} raw={raw_repr}"
                            )
                            continue  # skip empty/meaningless comments after cleaning

                        expanded = []
                        if classify and cleaned_content.strip():
                            print(f"    分类评论 [{ci+1}/{len(comments)}]...", end=" ", flush=True)
                            expanded = classify_content_expanded(cleaned_content)
                            tags_display = ", ".join(f"{e['tag']}-{e['sentiment']}" for e in expanded) if expanded else "/"
                            print(f"{tags_display}")
                        if expanded:
                            for item in expanded:
                                ws_comments.append([
                                    comment_counter, item["sentiment"], item["tag"],
                                    item["has_comparison"], cleaned_content, platform_name
                                ])
                        else:
                            ws_comments.append([comment_counter, "中性", "/", "否", cleaned_content, platform_name])
                        comment_counter += 1
                        written_in_url += 1

                    if len(comments):
                        logger.info(
                            f"[{platform_key}] 评论写入统计: raw={len(comments)} "
                            f"written={written_in_url} dropped={dropped_in_url}"
                        )

                    # Video processing
                    video_result = None
                    if video_processor and result.get("video_url"):
                        print(f"    处理视频...", end=" ", flush=True)
                        audio_url = result.get("audio_url")
                        video_result = video_processor.process_video(result["video_url"], audio_url, analyze=classify)
                        if video_result["status"] in ("success", "no_analysis"):
                            print(video_result["label"])
                            # Save audio base64 to file
                            audio_b64 = video_result.get("audio_base64")
                            if audio_b64:
                                try:
                                    audio_dir = config.OUTPUT_DIR / "audio_base64"
                                    audio_dir.mkdir(parents=True, exist_ok=True)
                                    serial = url.split("/")[-1].split("?")[0][:20]
                                    audio_file = audio_dir / f"{platform_key}_{serial}.txt"
                                    with open(audio_file, "w", encoding="utf-8") as af:
                                        af.write(audio_b64)
                                    print(f"    [音频base64] 已保存到 {audio_file.name}")
                                except Exception as ae:
                                    logger.warning(f"Failed to save audio base64: {ae}")
                        else:
                            print("跳过")

                    # Determine content display for Sheet 3
                    # transcription 来源：抖音走 video_processor；
                    # 视频号由 scrape() 经 parse_narration 自带 result["narration"]
                    # （OCR 评论后在线解析，替代录屏）
                    transcription = ""
                    video_summary = ""
                    if video_result and video_result["status"] in ("success", "no_analysis", "analysis_failed"):
                        transcription = video_result.get("transcription", "")
                        video_summary = video_result.get("summary", "")
                    if not transcription and result.get("narration"):
                        transcription = result["narration"]

                    parts = []
                    if post_content:
                        parts.append(f"正文：{post_content}")
                    if transcription:
                        parts.append(f"口播：{transcription}")
                    if parts:
                        content_display = "\n".join(parts)
                    elif video_result:
                        content_display = f"{video_result['label']}\n\n{video_summary}"
                    else:
                        content_display = ""

                    # Use transcription for tagging if available, fallback to summary/post
                    if transcription:
                        classify_text = transcription
                    elif video_result and video_result["status"] == "success":
                        classify_text = video_summary
                    else:
                        classify_text = post_content

                    # Write post content to Excel (one row per tag)
                    if content_display:
                        expanded = []
                        if classify and classify_text.strip():
                            print(f"    分类正文...", end=" ", flush=True)
                            expanded = classify_content_expanded(classify_text, is_post_content=True)
                            tags_display = ", ".join(f"{e['tag']}-{e['sentiment']}" for e in expanded) if expanded else "/"
                            print(f"{tags_display}")

                        if expanded:
                            for item in expanded:
                                row_data = [
                                    analysis_counter, row.get("title", ""), platform_name,
                                    content_display, item["sentiment"], item["tag"],
                                    item["has_comparison"], url
                                ]
                                ws_analysis.append(row_data)
                        else:
                            row_data = [
                                analysis_counter, row.get("title", ""), platform_name,
                                content_display, "/", "/", "否", url
                            ]
                            ws_analysis.append(row_data)

                        analysis_counter += 1

                    print(f"成功 ({count}条评论)")

                    results.append({
                        "platform": platform_key,
                        "title": row["title"],
                        "url": url,
                        "status": "成功",
                        "comment_count": count,
                        "error": "",
                    })
                    success_count += 1

                except Exception as e:
                    error_msg = classify_error(str(e))
                    is_no_comments = "暂无评论" in str(e)
                    print(f"{'跳过' if is_no_comments else '失败'} - {error_msg}")
                    if is_no_comments:
                        logger.info(f"[{platform_key}] {error_msg}: {url}")
                    else:
                        logger.error(f"[{platform_key}] 抓取失败: {url} - {e}")
                    results.append({
                        "platform": platform_key,
                        "title": row["title"],
                        "url": url,
                        "status": "失败",
                        "comment_count": 0,
                        "error": error_msg,
                    })
                    fail_count += 1

                    # Restart context if browser/context was closed
                    if "closed" in str(e).lower():
                        logger.info(f"[{platform_key}] Restarting browser context...")
                        try:
                            await scraper.stop()
                            _is_wx = platform_key in ("wechat", "wechat_channels")
                            _skw = {} if _is_wx else {"headless": headless}
                            await scraper.start(**_skw)
                        except Exception as re:
                            logger.error(f"[{platform_key}] Failed to restart: {re}")

                # ── Incremental checkpoint save ─────────────────
                try:
                    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
                    with open(checkpoint_path, "a", encoding="utf-8") as f:
                        f.write(_json.dumps(results[-1], ensure_ascii=False) + "\n")
                except Exception as e:
                    logger.warning(f"checkpoint 写入失败: {e}")

        finally:
            await scraper.stop()

    # Save results
    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(output_filepath)
        print(f"\n结果已保存: {output_filepath}")
    except PermissionError:
        alt_filepath = config.OUTPUT_DIR / f"scrape_result_{now:%Y%m%d_%H%M%S}_v2.xlsx"
        wb.save(alt_filepath)
        print(f"\n结果已保存: {alt_filepath}")
        print("(原文件被占用，使用了备用文件名)")

    write_summary_report(config.OUTPUT_DIR, results)

    # Save captured video URLs
    if video_urls_log:
        video_log_dir = config.OUTPUT_DIR / "video_urls"
        video_log_dir.mkdir(parents=True, exist_ok=True)
        video_log_path = video_log_dir / f"video_urls_{now:%Y%m%d_%H%M%S}.json"
        with open(video_log_path, "w", encoding="utf-8") as f:
            _json.dump(video_urls_log, f, ensure_ascii=False, indent=2)
        print(f"\n视频链接已保存: {video_log_path}")

    # ── Check if ALL URLs are done before cleaning checkpoint ────
    all_input_rows = read_input_excel(input_file)
    all_urls = {r["url"] for r in all_input_rows}
    all_scraped = completed_urls | {r["url"] for r in results}
    remaining = all_urls - all_scraped

    if not remaining:
        if checkpoint_path.exists():
            checkpoint_path.unlink()
            logger.info("所有链接已抓取，checkpoint 文件已清理")
    else:
        print(f"\n  剩余 {len(remaining)} 条链接未抓取，checkpoint 已保留。")
        print(f"  下次运行: python main.py --input {input_file} --resume --batch-size {batch_size}")

    print(f"\n{'=' * 50}")
    print(f"  本轮抓取完成！")
    print(f"  成功: {success_count}  失败: {fail_count}  总计: {success_count + fail_count}")
    if remaining:
        print(f"  剩余: {len(remaining)} 条")
    print(f"{'=' * 50}")

    if fail_count > 0:
        print(f"\n失败链接详情请查看: output/scrape_report_{now:%Y%m%d}.xlsx")


# ── CLI ──────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="多平台评论抓取工具（抖音/小红书/微博/头条/微信）",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
使用示例:
  python main.py --setup                     首次使用，自动安装环境
  python main.py --run --input input.xlsx    一键执行：自动登录+抓取（推荐）
  python main.py --run --input input.xlsx --resume --batch-size 20  带断点续抓
  python main.py --run --input input.xlsx --no-classify   只抓取不分类
  python main.py --run --input input.xlsx --platforms wechat,wechat_channels  只抓微信
  python main.py --login                     登录所有平台（打开浏览器，等待登录后保存）

微信公众号/视频号（PC 桌面模式）:
  1. 启动微信 PC 客户端并登录
  2. 运行: python main.py --run
  3. 默认使用 PC 相对坐标模式操作微信窗口
        """,
    )
    parser.add_argument("--setup", action="store_true", help="自动安装依赖环境（首次使用推荐）")
    parser.add_argument("--run", action="store_true", help="一键执行：自动检查登录+抓取（推荐）")
    parser.add_argument("--login", action="store_true", help="交互式登录各平台并保存Cookie")
    parser.add_argument("--login-open", action="store_true", help="打开登录浏览器（不阻塞，登录后需运行 --login-save）")
    parser.add_argument("--login-save", action="store_true", help="保存已打开浏览器的Cookie")
    parser.add_argument("--login-yuanbao", action="store_true",
                        help="单独登录腾讯元宝（视频号口播在线解析所需 cookie，存入 cookies/_yuanbao_profile）。"
                             "注：--login（全平台）已自动包含元宝登录，无需单独跑")
    parser.add_argument("--platforms", default="", help="指定平台(逗号分隔): douyin,xiaohongshu,weibo,toutiao")
    parser.add_argument("--input", default="", help="输入Excel文件路径")
    parser.add_argument("--resume", action="store_true", help="从上次中断处继续抓取（跳过已完成的URL）")
    parser.add_argument("--batch-size", type=int, default=0, help="每轮最多抓取N条（配合--resume避免超时）")
    parser.add_argument("--login-timeout", type=int, default=300, help="登录等待超时秒数（默认300）")
    parser.add_argument("--no-classify", action="store_true", help="跳过标签分类（只抓取不分析，加快速度）")
    parser.add_argument("--no-video", action="store_true", help="跳过视频口播处理（不下载视频、不转录）")
    parser.add_argument("--wechat-mode", default="pc", choices=["pc", "browser", "auto"],
                        help="微信操作模式: pc(桌面相对坐标,默认) browser(浏览器) auto(自动选择)")
    parser.add_argument("--wechat-calibrate", action="store_true",
                        help="校准微信PC桌面坐标（交互式点击记录UI元素位置）")
    parser.add_argument("--no-overlay", action="store_true",
                        help="校准时不显示屏幕叠加层（纯文字模式）")
    parser.add_argument("--headless", action="store_true",
                        help="无头模式运行浏览器（非微信平台适用）")
    args = parser.parse_args()

    if not args.setup:
        if not getattr(sys, 'frozen', False):
            missing = check_dependencies()
            if missing:
                print(f"[提示] 缺少依赖包: {', '.join(missing)}")
                print("正在自动安装...")
                if not install_dependencies():
                    sys.exit(1)

    if args.setup:
        if getattr(sys, 'frozen', False):
            print("[提示] 打包版本无需手动安装依赖，直接使用即可。")
            sys.exit(0)
        success = install_dependencies()
        sys.exit(0 if success else 1)

    setup_logging()
    import config
    config.COOKIE_DIR.mkdir(parents=True, exist_ok=True)
    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 落地内置默认微信坐标配置（冻结运行时首次启动自动写入 exe 同目录）。
    # 必须在版本推断 / 校准检查之前调用，保证 detect_version() 与
    # _ensure_wechat_config() 能看到落地后的默认配置。
    from scrapers.wechat import ensure_default_wechat_config
    ensure_default_wechat_config()

    # Handle --wechat-calibrate (before any run/login flow)
    if args.wechat_calibrate:
        ok = run_wechat_calibration(use_overlay=not args.no_overlay)
        sys.exit(0 if ok else 1)

    # Pre-check: if WeChat platforms are requested, ensure calibration exists
    if args.run or not args.login_open and not args.login_save and not args.login and not args.login_yuanbao:
        # Determine which platforms will be used
        if args.platforms:
            requested = set(p.strip() for p in args.platforms.split(",") if p.strip())
        else:
            requested = set()
        _ensure_wechat_config(requested)

    if getattr(sys, 'frozen', False):
        print("=" * 60)
        print("  评论抓取工具 v1.0")
        print("=" * 60)
        if not config.MIMO_API_KEY:
            print("  [提示] 未配置 MiMo API Key，分类功能将被跳过。")
            print("  配置方法: 编辑同目录下的 config.ini 文件")
        print()

    # 裸运行（无动作参数）：显示启动菜单，提供抓取/校准/退出入口。
    # 放在输入文件检查之前，保证没有输入 Excel 时也能进入校准模式。
    bare_run = not (args.run or args.login or args.login_open or
                    args.login_save or args.login_yuanbao)
    if bare_run:
        if not _show_startup_menu(args):
            sys.exit(0)

    input_file = args.input if args.input else str(config.INPUT_FILE)
    if not Path(input_file).exists():
        print(f"[错误] 找不到输入文件: {input_file}")
        print(f"请将 .xlsx 文件放到 input/ 目录下")
        sys.exit(1)

    all_platforms = ["douyin", "xiaohongshu", "weibo", "toutiao", "bilibili"]

    if args.run:
        platforms_filter = None
        if args.platforms:
            platforms_filter = [p.strip() for p in args.platforms.split(",") if p.strip()]
        asyncio.run(run_all(input_file, platforms_filter, resume=args.resume,
                            batch_size=args.batch_size, login_timeout=args.login_timeout,
                            classify=not args.no_classify, process_video=not args.no_video,
                            headless=args.headless))
    elif args.login_open:
        if args.platforms:
            platform_list = [p.strip() for p in args.platforms.split(",") if p.strip()]
        else:
            platform_list = all_platforms
        asyncio.run(login_open_all(platform_list))
    elif args.login_save:
        asyncio.run(login_save_all(all_platforms))
    elif args.login:
        if args.platforms:
            platform_list = [p.strip() for p in args.platforms.split(",") if p.strip()]
        else:
            platform_list = all_platforms
        asyncio.run(login_open_all(platform_list))
        # 元宝登录与平台登录一并完成：视频号口播在线解析所需 cookie 存入
        # cookies/_yuanbao_profile（持久化 profile，不走 config.ini）。
        # 默认全平台、或 --platforms 含 wechat/wechat_channels 时触发。
        _want_yuanbao = (not args.platforms) or any(
            p in platform_list for p in ("wechat", "wechat_channels")
        )
        if _want_yuanbao:
            from scrapers.wechat.channels_sph_parser import ChannelsSphParser
            print("\n" + "=" * 56)
            print("  接下来登录腾讯元宝（视频号口播在线解析所需）")
            print("  扫码登录后自动检测并保存，或回终端按 Enter 确认")
            print("=" * 56)
            asyncio.run(ChannelsSphParser.login_interactive(wait_seconds=args.login_timeout))
    elif args.login_yuanbao:
        from scrapers.wechat.channels_sph_parser import ChannelsSphParser
        ok = asyncio.run(ChannelsSphParser.login_interactive(wait_seconds=args.login_timeout))
        sys.exit(0 if ok else 1)
    else:
        # 裸运行到达这里：启动菜单已确认开始抓取
        platforms_filter = None
        if args.platforms:
            platforms_filter = [p.strip() for p in args.platforms.split(",") if p.strip()]
        asyncio.run(scrape_all(input_file, platforms_filter, resume=args.resume,
                               batch_size=args.batch_size, classify=not args.no_classify,
                               process_video=not args.no_video, headless=args.headless))


if __name__ == "__main__":
    main()
